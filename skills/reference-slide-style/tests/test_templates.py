"""Regression checks for the source-only bilingual template distribution."""
from pathlib import Path
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation

SKILL = Path(__file__).resolve().parents[1]


def slide_text(slide):
    values = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip():
            values.append(shape.text)
        if shape.has_table:
            values.extend(cell.text for row in shape.table.rows for cell in row.cells if cell.text.strip())
    return values


class TemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.workspace = TemporaryDirectory(prefix="reference-style-test-")
        cls.addClassCleanup(cls.workspace.cleanup)
        cls.root = Path(cls.workspace.name)
        cls.skill = cls.root / "skill"
        shutil.copytree(SKILL, cls.skill, ignore=shutil.ignore_patterns(
            "*.pptx", "*.potx", "__pycache__", ".venv", "previews"))
        cls.output = cls.root / "output"
        subprocess.run([sys.executable, str(cls.skill / "scripts/build_templates.py"),
                        "--output-dir", str(cls.output)], cwd=cls.root, check=True,
                       capture_output=True, text=True)

    def run_checker(self, paths, guide=False):
        cmd = [sys.executable, str(self.skill / "scripts/check_templates.py")]
        cmd.extend(str(path) for path in paths)
        if guide:
            cmd.append("--guide")
        result = subprocess.run(cmd, cwd=self.root, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(result.stdout.count(": PASS;"), len(paths))

    def test_six_office_outputs(self):
        expected = {f"{kind}-{lang}.{ext}" for lang in ("en", "zh-CN")
                    for kind, ext in (("blank-layouts", "pptx"),
                                      ("layout-guide", "pptx"),
                                      ("reference-style", "potx"))}
        actual = {p.name for p in self.output.iterdir() if p.suffix in (".pptx", ".potx")}
        self.assertEqual(actual, expected)

    def test_blank_and_potx_structural_checks(self):
        paths = sorted(self.output.glob("blank-layouts-*.pptx"))
        paths += sorted(self.output.glob("reference-style-*.potx"))
        self.run_checker(paths)

    def test_guide_structural_checks(self):
        self.run_checker(sorted(self.output.glob("layout-guide-*.pptx")), guide=True)

    def test_blank_content_and_closing_copy(self):
        for lang, thanks in (("en", ["Thank you", "Questions & discussion"]),
                             ("zh-CN", ["谢谢聆听", "欢迎交流与提问"])):
            with self.subTest(language=lang):
                deck = Presentation(self.output / f"blank-layouts-{lang}.pptx")
                self.assertEqual(len(deck.slides), 14)
                for index, slide in enumerate(deck.slides):
                    self.assertEqual(slide_text(slide), thanks if index == 13 else [])

    def test_table_and_geometry_match_across_languages(self):
        decks = [Presentation(self.output / f"blank-layouts-{lang}.pptx")
                 for lang in ("en", "zh-CN")]
        for deck in decks:
            tables = [shape.table for shape in deck.slides[7].shapes if shape.has_table]
            self.assertEqual(len(tables), 1)
            self.assertEqual((len(tables[0].columns), len(tables[0].rows)), (5, 11))
            self.assertEqual(len(deck.slide_layouts), 14)
        for en, zh in zip(decks[0].slides, decks[1].slides):
            geometry = lambda slide: [(s.name, s.left, s.top, s.width, s.height)
                                      for s in slide.shapes]
            self.assertEqual(geometry(en), geometry(zh))

    def test_no_binary_seed_is_required_or_written(self):
        self.assertFalse((self.skill / "assets/starter-layouts.pptx").exists())
        self.assertFalse(list(self.skill.rglob("*.pptx")))
        self.assertFalse(list(self.skill.rglob("*.potx")))


if __name__ == "__main__":
    unittest.main()
