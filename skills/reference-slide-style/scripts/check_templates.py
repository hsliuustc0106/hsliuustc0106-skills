#!/usr/bin/env python3
"""Check v2 template structure and deliberately blank slide content.

This is not a renderer, an OOXML schema validator, or a factual-content checker.
POTX content type is checked directly; a temporary in-memory PPTX view allows
python-pptx to inspect template geometry without changing the original file.
"""
from __future__ import annotations
import argparse
from io import BytesIO
from pathlib import Path
import sys
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree
from pptx import Presentation

P='http://schemas.openxmlformats.org/presentationml/2006/main'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
PPTX='application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml'
POTX='application/vnd.openxmlformats-officedocument.presentationml.template.main+xml'


def load(path):
    errors=[]
    with ZipFile(path) as z:
        bad=z.testzip()
        if bad:errors.append('ZIP integrity failure: '+bad)
        types=etree.fromstring(z.read('[Content_Types].xml'))
        main=next((e for e in types if e.get('PartName')=='/ppt/presentation.xml'),None)
        expected=POTX if path.suffix.lower()=='.potx' else PPTX
        if main is None or main.get('ContentType')!=expected:errors.append('Unexpected presentation content type.')
        for name in z.namelist():
            if name.lower().endswith(('.ttf','.otf','.ttc','.fntdata','.odttf')):errors.append('Font binary included: '+name)
            if name.startswith(('ppt/slides/slide','ppt/slideLayouts/slideLayout')) and name.endswith('.xml'):
                root=etree.fromstring(z.read(name))
                ids=root.xpath('.//p:cNvPr/@id',namespaces={'p':P})
                if len(ids)!=len(set(ids)):errors.append('Duplicate shape IDs in '+name)
                phs=root.xpath('.//p:ph/@idx',namespaces={'p':P})
                if len(phs)!=len(set(phs)):errors.append('Duplicate placeholder indices in '+name)
                if root.xpath('.//a:outerShdw | .//a:innerShdw',namespaces={'a':A}):errors.append('Explicit shadow in '+name)
        buf=BytesIO()
        with ZipFile(buf,'w',ZIP_DEFLATED) as dst:
            for name in z.namelist():
                data=z.read(name)
                if name=='[Content_Types].xml' and path.suffix.lower()=='.potx':
                    main.set('ContentType',PPTX)
                    data=etree.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
                dst.writestr(name,data)
    buf.seek(0)
    return Presentation(buf),errors


def check(path,guide=False):
    prs,errors=load(path)
    if len(prs.slides)!=14:errors.append(f'Expected 14 slides; got {len(prs.slides)}.')
    if len(prs.slide_layouts)!=14:errors.append(f'Expected 14 custom layouts; got {len(prs.slide_layouts)}.')
    if abs(prs.slide_width/914400-13.333333333)>1e-5 or abs(prs.slide_height/914400-7.5)>1e-5:errors.append('Canvas must be 16:9 at 13.333333 x 7.5 inches.')
    texts=[];table_count=0;ph_count=0;static_count=0
    for group,collection in [('slide',prs.slides),('layout',prs.slide_layouts)]:
        for i,s in enumerate(collection,1):
            for sh in s.shapes:
                if group=='layout' and not sh.is_placeholder:static_count+=1
                if group=='slide' and sh.is_placeholder:ph_count+=1
                try:
                    x,y,w,h=sh.left,sh.top,sh.width,sh.height
                    if None not in (x,y,w,h) and (x<-50 or y<-50 or x+w>prs.slide_width+50 or y+h>prs.slide_height+50):errors.append(f'{group} {i}: out-of-bounds {sh.name}.')
                except (AttributeError,TypeError):pass
                if group=='slide' and sh.has_text_frame and sh.text.strip():
                    texts.append((i,sh.name,sh.text.strip()))
                    if not guide and (i!=14 or sh.name not in ('THANKS_TITLE','THANKS_SUBTITLE')):errors.append(f'Unexpected content in blank slide {i}: {sh.name}.')
                if group=='slide' and sh.has_table:
                    table_count+=1
                    tb=sh.table
                    if len(tb.columns)!=5 or len(tb.rows)!=11:errors.append('Prepared table must have five columns and eleven total rows.')
                    if abs(sum(c.width for c in tb.columns)-sh.width)>20 or abs(sum(r.height for r in tb.rows)-sh.height)>20:errors.append('Native table frame and grid disagree.')
                    for row in tb.rows:
                        for cell in row.cells:
                            if not guide and cell.text.strip():errors.append('Unexpected content in blank table.')
    if table_count!=1:errors.append(f'Expected one native table; got {table_count}.')
    if not guide:
        zh='zh-CN' in path.name
        expect=('谢谢聆听','欢迎交流与提问') if zh else ('Thank you','Questions & discussion')
        ending=tuple(txt for i,name,txt in texts if name in ('THANKS_TITLE','THANKS_SUBTITLE'))
        if ending!=expect:errors.append('Wrong or missing closing text.')
        if len(texts)!=2:errors.append(f'Expected exactly two visible text fields; got {len(texts)}.')
    print(f'{path.name}: {"FAIL" if errors else "PASS"}; slides={len(prs.slides)}, layouts={len(prs.slide_layouts)}, placeholders={ph_count}, fixed-layout-shapes={static_count}, native-tables={table_count}')
    for msg in errors:print('  - '+msg)
    return not errors


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('files',nargs='+',type=Path)
    p.add_argument('--guide',action='store_true',help='Allow field-label text in a labelled guide.')
    a=p.parse_args();ok=True
    for path in a.files:
        try:ok=check(path,a.guide) and ok
        except (ValueError,KeyError,OSError) as e:print(f'{path}: ERROR {e}');ok=False
    sys.exit(0 if ok else 1)

if __name__=='__main__':main()
