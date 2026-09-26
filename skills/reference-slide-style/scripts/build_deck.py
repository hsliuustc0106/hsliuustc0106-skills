#!/usr/bin/env python3
"""Build editable slides from the measured reference style.

Tested with Python 3.13.5 and python-pptx 1.0.2. No network access required.
Run from any directory; paths default to this skill package.
All sample text is instructional placeholder content, not research evidence.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Sequence
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TOKENS = json.loads((ROOT / 'assets' / 'design-tokens.json').read_text(encoding='utf-8'))
C = TOKENS['colors']
T = TOKENS['typography']
L = TOKENS['layouts']


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(C.get(value, value).lstrip('#'))


def paint(shape: Any, fill: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    # Override the default Office shape effect so every card remains flat.
    sp = shape._element.spPr
    if sp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst') is None:
        sp.append(OxmlElement('a:effectLst'))
    for ref in shape._element.xpath('./p:style/a:effectRef'):
        ref.set('idx', '0')


def box(slide: Any, name: str, rect: Sequence[float], fill: str,
        rounded: bool = False, adjustment: float = 0.12) -> Any:
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, *(Inches(v) for v in rect))
    sh.name = name
    paint(sh, fill)
    if rounded and len(sh.adjustments):
        sh.adjustments[0] = adjustment
    return sh


def text(slide: Any, name: str, rect: Sequence[float], content: str,
         size: float, color: str = 'navy', bold: bool = False,
         align: str = 'left', anchor: str = 'top', spacing: float = 1.12,
         font: str = 'Arial') -> Any:
    sh = slide.shapes.add_textbox(*(Inches(v) for v in rect))
    sh.name = name
    tf = sh.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if anchor == 'middle' else MSO_ANCHOR.TOP
    for i, line in enumerate(str(content).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        # Explicit CJK fallback is a portability choice, not an original font claim.
        ea = OxmlElement('a:ea')
        ea.set('typeface', T['cjk_fallback_recommended'])
        r._r.get_or_add_rPr().append(ea)
    return sh


def arrow(slide: Any, name: str, start: Sequence[float], end: Sequence[float],
          color: str = 'blue', head: bool = True) -> Any:
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
        Inches(start[0]), Inches(start[1]), Inches(end[0]), Inches(end[1]))
    sh.name = name
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(TOKENS['components']['connector_width_pt'])
    if head:
        tail = OxmlElement('a:tailEnd')
        tail.set('type', 'triangle')
        tail.set('w', 'sm')
        tail.set('len', 'sm')
        sh._element.spPr.get_or_add_ln().append(tail)
    return sh


def source_notes(slide: Any, data: dict[str, Any], source: str) -> None:
    slide.notes_slide.notes_text_frame.text = (
        f"LAYOUT ORIGIN: {source}\n"
        "STYLE SOURCE: ai-application-innovation-20-pages.pptx (user-provided reference).\n"
        "This starter uses instructional placeholder text; it makes no research or market claims.\n"
        "Replace the title with one conclusion; state the audience, scope, or caveat in the subtitle.\n"
        "Update branding, source dates, citations, and page numbers before delivery.\n"
        "Keep content editable. Split an overfull slide instead of shrinking essential text.\n"
        + str(data.get('speaker_notes', ''))
    )


def chrome(prs: Any, meta: dict[str, Any], data: dict[str, Any], index: int,
           cover: bool = False) -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb('white')
    g = TOKENS['chrome']
    if not cover:
        text(slide, 'HEADLINE', g['headline'], data['title'], T['headline'], 'blue', True)
        text(slide, 'SUBTITLE', g['subtitle'], data.get('subtitle', ''), T['subtitle'], 'slate')
        box(slide, 'HEADER_RULE', g['header_rule'], 'rule')
    box(slide, 'FOOTER_RULE', g['footer_rule'], 'rule')
    section = data.get('section', data['layout'].replace('_', ' ').title())
    text(slide, 'FOOTER_LABEL', g['footer_label'],
         f"{meta.get('deck_label', '[DECK]')}  /  {section}", T['footer'], 'slate')
    text(slide, 'BRAND_PLACEHOLDER', g['brand'], meta.get('brand', '[BRAND]'),
         12, 'purple', True, 'right')
    text(slide, 'PAGE_NUMBER', g['page'], f'{index:02d}', T['footer'], 'slate', align='right')
    if data.get('source_note'):
        text(slide, 'SOURCE_NOTE', g['source_note'], data['source_note'], T['source_note'], 'slate')
    return slide


def takeaway(slide: Any, value: str) -> None:
    if not value:
        return
    box(slide, 'TAKEAWAY_BACKGROUND', TOKENS['components']['takeaway_bg'], 'pale_blue', True)
    text(slide, 'TAKEAWAY', TOKENS['components']['takeaway_text'], value, T['takeaway'], 'blue', True)


def build_cover(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n, True)
    a = L['cover']
    box(s, 'HERO_REPLACE_WITH_IMAGE', a['hero'], 'pale_blue')
    text(s, 'HERO_PLACEHOLDER_LABEL', [9.0, 2.25, 3.5, 1.15],
         'COVER IMAGE\nReplace with an approved visual', 17, 'purple', align='center', anchor='middle')
    box(s, 'COVER_WHITE_PANEL', a['white_panel'], 'white')
    text(s, 'COVER_KICKER', a['kicker'], d.get('kicker', '[SUBJECT]'), T['cover_kicker'], 'blue', True)
    text(s, 'COVER_TITLE', a['title'], d['title'], T['cover_headline'], 'navy', True)
    text(s, 'COVER_SUBTITLE', a['subtitle'], d.get('subtitle', ''), T['cover_subtitle'], 'slate')
    text(s, 'COVER_META', a['meta_safe'], d.get('meta', ''), T['cover_meta'], 'blue', True)
    source_notes(s, d, 'Slide 1. Image and original wordmark intentionally not carried over. Cover meta box enlarged for two lines.')


def build_agenda(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['agenda']
    for i, item in enumerate(d['items']):
        y = a['row_y'][i]
        box(s, f'AGENDA_ROW_{i+1}', [a['row_x'], y, a['row_w'], a['row_h']],
            ['pale_blue', 'pale_purple', 'pale_teal'][i], True)
        text(s, f'AGENDA_NUMBER_{i+1}', [0.86, y+.28, .65, .45], item['number'], 25, 'blue', True)
        text(s, f'AGENDA_TITLE_{i+1}', [a['title_x'], y+.31, a['title_w'], .44], item['title'], 23, 'navy', True)
        text(s, f'AGENDA_DETAIL_{i+1}', [a['detail_x'], y+.21, a['detail_w'], .73], item['detail'], 19, 'slate')
    takeaway(s, d.get('takeaway', ''))
    source_notes(s, d, 'Slide 2. Detail boxes have extra vertical room for two English lines.')


def build_overview(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['overview']
    for i, item in enumerate(d['cards']):
        x, y = a['x'][i % 3], a['y'][i // 3]
        box(s, f'OVERVIEW_CARD_{i+1}', [x, y, a['card_w'], a['card_h']], 'pale_blue', True)
        text(s, f'CARD_NUMBER_{i+1}', [x+.23, y+.18, .6, .4], item.get('number', f'{i+1:02d}'), T['grid_number'], 'blue', True)
        text(s, f'CARD_TITLE_{i+1}', [x+.23, y+.74, 3.39, .39], item['title'], T['grid_title'], 'navy', True)
        text(s, f'CARD_BODY_{i+1}', [x+.23, y+1.22, 3.39, .6], item['body'], T['grid_body'], 'slate')
    source_notes(s, d, 'Slides 3 and 15. Six parallel units; keep body copy to two short lines.')


def build_triptych(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['triptych']
    for i, item in enumerate(d['cards']):
        x, y = a['x'][i], a['y']
        accent = ['blue', 'purple', 'teal'][i]
        box(s, f'COMPARISON_CARD_{i+1}', [x, y, a['w'], a['h']], 'pale_blue', True)
        box(s, f'CARD_ACCENT_{i+1}', [x, y, a['w'], .045], accent)
        text(s, f'CARD_LABEL_{i+1}', [x+.25, y+.25, 3.35, .3], item['label'], T['card_label'], accent, True)
        text(s, f'CARD_TITLE_{i+1}', [x+.25, y+.76, 3.35, .8], item['title'], T['card_title'], 'navy', True)
        text(s, f'CARD_BODY_{i+1}', [x+.25, y+1.658, 3.35, 1.639], item['body'], T['card_body'], 'slate')
        text(s, f'CARD_NOTE_{i+1}', [x+.25, y+3.43, 3.35, .33], item['note'], T['card_note'], accent, True)
    takeaway(s, d.get('takeaway', ''))
    origin = 'Slide 19 (roadmap)' if d['layout'] == 'roadmap' else 'Slides 4, 6, 8, 10 and 20 (evidence / decision triptych)'
    source_notes(s, d, origin + '. Keep the bottom note to one line; split complex evidence into a follow-up slide.')


def build_architecture(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['architecture']
    accs = ['blue', 'purple', 'teal', 'blue', 'purple', 'teal', 'blue']
    fills = ['pale_blue', 'pale_purple', 'pale_teal', 'pale_blue', 'pale_purple', 'pale_teal', 'pale_purple']
    for i, value in enumerate(d['steps']):
        x = a['step_x'] + i * (a['step_w'] + a['step_gap'])
        box(s, f'WORKFLOW_STEP_{i+1}', [x, a['step_y'], a['step_w'], a['step_h']], fills[i], True)
        text(s, f'WORKFLOW_TEXT_{i+1}', [x+.014, a['step_y']+.014, a['step_w']-.028, .472], value,
             T['workflow_step'], accs[i], True, 'center', 'middle', 1.0)
    box(s, 'ARCHITECTURE_CONTAINER', a['container'], 'pale_blue', True)
    text(s, 'SYSTEM_TITLE', a['system_title'], d['system_title'], T['architecture_title'], 'blue', True, 'center', 'middle', 1.0)
    layer_accents = ['blue', 'purple', 'blue', 'teal']
    for i, layer in enumerate(d['layers']):
        y = a['layer_y'][i]
        box(s, f'LAYER_{i+1}', [a['layer_x'], y, a['layer_w'], a['layer_h']], 'white', True)
        text(s, f'LAYER_TITLE_{i+1}', [.794, y+.052, 5.692, .25], layer['title'], T['layer_title'], layer_accents[i], True, 'center', 'middle', 1.0)
        text(s, f'LAYER_BODY_{i+1}', [.794, y+.321, 5.692, .23], layer['body'], T['layer_body'], 'navy', False, 'center', 'middle', 1.0)
        if i < 3:
            arrow(s, f'LAYER_ARROW_{i+1}', [3.639, y+.598], [3.639, a['layer_y'][i+1]-.005])
    box(s, 'PLATFORM_BAND', a['platform_band'], 'pale_purple', True)
    text(s, 'PLATFORM_BAND_TEXT', [.656, 6.024, 5.969, .326], d['platform_label'], 16, 'purple', True, 'center', 'middle', 1.0)
    box(s, 'DEPENDENCY_BAND', a['dependency_band'], 'pale_teal')
    text(s, 'DEPENDENCY_TEXT', [.656, 6.413, 5.969, .3], d['dependency_label'], 14, 'teal', True, 'center', 'middle', 1.0)
    for i, callout in enumerate(d['callouts']):
        accent = layer_accents[i]
        y = a['callout_title_y'][i]
        box(s, f'CALLOUT_BAR_{i+1}', [a['callout_bar_x'], y+.042, .05, .361], accent)
        text(s, f'CALLOUT_TITLE_{i+1}', [a['callout_x'], y, a['callout_w'], .417], callout['title'], T['callout_title'], accent, True, spacing=1.0)
        text(s, f'CALLOUT_BODY_{i+1}', [a['callout_x'], a['callout_body_y'][i], a['callout_w'], .653], callout['body'], T['callout_body'], 'slate', spacing=1.0)
    source_notes(s, d, 'Slides 5, 7, 9, 11, 12 and 14. Seven short steps, four layers, four callouts. Source note remains above the footer.')


def cell_border(cell: Any, color: str = 'rule') -> None:
    pr = cell._tc.get_or_add_tcPr()
    for edge in ['L', 'R', 'T', 'B']:
        ln = OxmlElement(f'a:ln{edge}')
        ln.set('w', '9525')
        fill = OxmlElement('a:solidFill')
        clr = OxmlElement('a:srgbClr')
        clr.set('val', C[color])
        fill.append(clr)
        ln.append(fill)
        pr.append(ln)


def build_table(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['table']
    widths = a['column_widths']
    rows = [d['headers']] + d['rows']
    sh = s.shapes.add_table(len(rows), 5, Inches(a['x']), Inches(a['y']),
                           Inches(sum(widths)), Inches(.42*len(rows)))
    sh.name = 'EVIDENCE_TABLE'
    tbl = sh.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, width in enumerate(widths):
        tbl.columns[j].width = Inches(width)
    highlights = {tuple(v) for v in d.get('highlight_cells', [])}
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(.42)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Inches(.10)
            cell.margin_top = cell.margin_bottom = Inches(.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            fill = 'blue' if i == 0 else ('pale_teal' if (i-1, j) in highlights else ('pale_blue' if i % 2 else 'white'))
            cell.fill.fore_color.rgb = rgb(fill)
            cell_border(cell)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(T['table'])
                    r.font.bold = i == 0
                    r.font.color.rgb = rgb('white' if i == 0 else 'navy')
    source_notes(s, d, 'Slide 13. Original table has inconsistent OOXML frame and grid dimensions. This version uses coherent 0.42-inch rows, reduced vertical cell padding, and a native editable table.')


def build_platform(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['platform']
    accents = ['blue', 'blue', 'purple', 'teal']
    fills = ['pale_blue', 'pale_blue', 'pale_purple', 'pale_teal']
    for i, item in enumerate(d['apps']):
        x = a['app_x'][i]
        box(s, f'APPLICATION_{i+1}', [x, a['app_y'], a['app_w'], a['app_h']], fills[i], True)
        text(s, f'APPLICATION_TITLE_{i+1}', [x+.07, 2.194, 2.707, .528], item['title'], 20, accents[i], True, 'center', 'middle', 1.0)
        text(s, f'APPLICATION_BODY_{i+1}', [x+.07, 2.764, 2.707, .417], item['body'], 16.5, 'navy', align='center', anchor='middle', spacing=1.0)
        arrow(s, f'PLATFORM_ARROW_{i+1}', [x+a['app_w']/2, 3.791], [x+a['app_w']/2, 3.361], accents[i])
    box(s, 'SHARED_FOUNDATION', a['foundation'], 'pale_purple', True)
    text(s, 'FOUNDATION_TITLE', [.831, 3.861, 11.59, .639], d['platform_name'], 25, 'purple', True, 'center', 'middle', 1.0)
    for i, value in enumerate(d['capabilities']):
        x = a['capability_x'][i]
        box(s, f'CAPABILITY_{i+1}', [x, 4.625, 2.208, .597], 'white', True)
        text(s, f'CAPABILITY_TEXT_{i+1}', [x+.10, 4.70, 2.008, .44], value, 17, 'navy', align='center', anchor='middle', spacing=1.0)
    text(s, 'FOUNDATION_NOTE', [.831, 5.319, 11.59, .417], d['platform_note'], 16, 'purple', align='center', anchor='middle', spacing=1.0)
    box(s, 'INPUT_OUTPUT_BAND', a['io_band'], 'pale_teal', True)
    text(s, 'INPUT_OUTPUT_TEXT', [.739, 6.029, 11.772, .607], d['io_text'], 16, 'teal', True, 'center', 'middle', 1.0)
    source_notes(s, d, 'Slide 16. Four applications above a shared foundation; arrows point from the foundation to the applications.')


def build_process(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['process']
    for i, item in enumerate(d['cards']):
        x, y = a['x'][i], a['y']
        box(s, f'PROCESS_CARD_{i+1}', [x, y, a['w'], a['h']], ['pale_blue','pale_purple','pale_teal'][i], True)
        box(s, f'PROCESS_ACCENT_{i+1}', [x, y, a['w'], .045], 'blue')
        text(s, f'PROCESS_LABEL_{i+1}', [x+.25, y+.25, 3.35, .3], item['label'], 12, 'blue', True)
        text(s, f'PROCESS_TITLE_{i+1}', [x+.25, y+.76, 3.35, .8], item['title'], 24, 'navy', True)
        text(s, f'PROCESS_BODY_{i+1}', [x+.25, y+1.71, 3.35, 1.01], item['body'], 19, 'slate')
        if i < 2:
            arrow(s, f'PROCESS_ARROW_{i+1}', [x+3.89, 3.59], [x+4.09, 3.59])
    text(s, 'PROCESS_SEQUENCE', a['sequence'], d['sequence'], 20, 'navy', True)
    takeaway(s, d.get('takeaway', ''))
    source_notes(s, d, 'Slide 17. Three stages with arrows; the sequence and takeaway have distinct roles.')


def build_feedback(prs: Any, meta: dict[str, Any], d: dict[str, Any], n: int) -> None:
    s = chrome(prs, meta, d, n)
    a = L['feedback']
    box(s, 'FEEDBACK_INPUT_PANEL', a['input_panel'], 'pale_purple', True, .16667)
    text(s, 'FEEDBACK_INPUT_TITLE', [.85, 2.20, 6.9, .34], d['input_title'], 21, 'purple', True, spacing=1.06)
    for i, item in enumerate(d['inputs']):
        x = .77 + i*1.77
        box(s, f'INPUT_LABEL_BG_{i+1}', [x, 2.59, 1.64, .32], 'white', True)
        text(s, f'INPUT_LABEL_{i+1}', [x+.025, 2.60, 1.59, .30], item['title'], 13.1, 'navy', True, 'center', 'middle', 1.0)
        text(s, f'INPUT_DETAIL_{i+1}', [x, 2.96, 1.64, .27], item['body'], 10.5, 'slate', align='center', spacing=1.0)
    box(s, 'FEEDBACK_LEDGER', a['ledger'], 'pale_blue', True, .16667)
    text(s, 'FEEDBACK_LEDGER_TEXT', [.75, 3.55, 7.03, .52], d['ledger'], 15.4, 'blue', True, 'center', 'middle', 1.06)
    arrow(s, 'INPUT_TO_LEDGER', [4.2, 3.29], [4.2, 3.50])
    for i, item in enumerate(d['stages']):
        x,y,w,h = a['stage_boxes'][i]
        accent = 'purple' if i == 2 else 'blue'
        box(s, f'FEEDBACK_STAGE_{i+1}', [x,y,w,h], 'pale_purple' if i==2 else 'pale_blue', True, .16667)
        text(s, f'FEEDBACK_STAGE_TITLE_{i+1}', [x+.06, y+.06, w-.12, .31], item['title'], 15.5, accent, True, 'center', 'middle', 1.0)
        text(s, f'FEEDBACK_STAGE_BODY_{i+1}', [x+.09, y+.47, w-.18, .58], item['body'], 14.5, 'navy', align='center', anchor='middle', spacing=1.06)
    arrow(s, 'LEDGER_TO_TRIAGE', [1.6,4.12], [1.6,4.36])
    arrow(s, 'TRIAGE_TO_OWNER', [2.63,4.90], [2.89,4.90])
    arrow(s, 'OWNER_TO_RETEST', [4.99,4.90], [5.24,4.90])
    box(s, 'FEEDBACK_RETURN', a['return_band'], 'pale_teal', True, .16667)
    text(s, 'FEEDBACK_RETURN_TEXT', [.78, 5.78, 6.97, .43], d['return_text'], 18, 'teal', True, 'center', 'middle', 1.06)
    text(s, 'INTERNAL_TITLE', a['internal_title'], d['internal_title'], 21, 'blue', True)
    text(s, 'INTERNAL_BODY', a['internal_body'], d['internal_body'], 16.5, 'navy', spacing=1.06)
    text(s, 'EXTERNAL_TITLE', a['external_title'], d['external_title'], 21, 'purple', True)
    text(s, 'EXTERNAL_BODY', a['external_body'], d['external_body'], 16.5, 'slate', spacing=1.06)
    box(s, 'IMPROVEMENT_BAND', a['improvement'], 'pale_purple', True)
    text(s, 'IMPROVEMENT_TEXT', [8.40,6.29,4.10,.39], d['improvement'], 14.5, 'purple', True, 'center', 'middle', 1.0)
    arrow(s, 'LOOP_OUT', [7.94,6.0], [8.1,6.0], head=False)
    arrow(s, 'LOOP_UP', [8.1,6.0], [8.1,2.34], head=False)
    arrow(s, 'LOOP_IN', [8.1,2.34], [7.94,2.34])
    source_notes(s, d, 'Slide 18. A traceable feedback loop on the left; internal and external validation on the right. Compact labels shortened to prevent inherited word breaks.')


BUILDERS = {
    'cover': build_cover, 'agenda': build_agenda, 'overview': build_overview,
    'comparison': build_triptych, 'decision': build_triptych, 'roadmap': build_triptych,
    'architecture': build_architecture, 'table': build_table,
    'platform': build_platform, 'process': build_process, 'feedback': build_feedback,
}


def validate_content(payload: dict[str, Any]) -> None:
    slides = payload.get('slides')
    if not isinstance(slides, list) or not slides:
        raise ValueError('Content must contain a non-empty slides array.')
    for i, d in enumerate(slides, 1):
        if d.get('layout') not in BUILDERS:
            raise ValueError(f"Slide {i}: unknown layout {d.get('layout')!r}. Choose {', '.join(BUILDERS)}.")
        if not d.get('title'):
            raise ValueError(f'Slide {i}: title is required.')
        counts = {
            'agenda': {'items': 3}, 'overview': {'cards': 6},
            'comparison': {'cards': 3}, 'decision': {'cards': 3}, 'roadmap': {'cards': 3},
            'architecture': {'steps': 7, 'layers': 4, 'callouts': 4},
            'platform': {'apps': 4, 'capabilities': 5}, 'process': {'cards': 3},
            'feedback': {'inputs': 4, 'stages': 3},
        }.get(d['layout'], {})
        for key, expected in counts.items():
            if len(d.get(key, [])) != expected:
                raise ValueError(f'Slide {i}: {key} must contain exactly {expected} entries for this starter geometry.')
        if d['layout'] == 'table':
            if len(d.get('headers', [])) != 5 or not 1 <= len(d.get('rows', [])) <= 10:
                raise ValueError(f'Slide {i}: table needs 5 headers and 1-10 data rows.')
            if any(len(row) != 5 for row in d['rows']):
                raise ValueError(f'Slide {i}: each table row needs 5 cells.')


def set_theme(prs: Any) -> None:
    """Set the actual visual palette, not the reference's unrelated Office defaults."""
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype.endswith('/theme'):
            part = rel.target_part
            root = etree.fromstring(part.blob)
            mapping = {'dk1':'navy', 'lt1':'white', 'dk2':'slate', 'lt2':'pale_blue',
                       'accent1':'blue', 'accent2':'purple', 'accent3':'teal',
                       'accent4':'pale_blue', 'accent5':'pale_purple', 'accent6':'pale_teal',
                       'hlink':'blue', 'folHlink':'purple'}
            scheme = root.find('.//a:clrScheme', ns)
            scheme.set('name', 'Reference Slide Style')
            for key, token in mapping.items():
                elem = scheme.find('a:'+key, ns)
                for child in list(elem):
                    elem.remove(child)
                clr = etree.SubElement(elem, '{'+ns['a']+'}srgbClr')
                clr.set('val', C[token])
            # Some renderers still inherit theme effects despite an empty local list.
            for effect in root.findall('.//a:effectStyleLst/a:effectStyle', ns):
                for child in list(effect):
                    effect.remove(child)
                etree.SubElement(effect, '{'+ns['a']+'}effectLst')
            for effect_ref in root.findall('.//a:effectRef', ns):
                effect_ref.set('idx', '0')
            for latin in root.findall('.//a:fontScheme//a:latin', ns):
                latin.set('typeface', 'Arial')
            for ea in root.findall('.//a:fontScheme//a:ea', ns):
                ea.set('typeface', T['cjk_fallback_recommended'])
            part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def make_deck(content_path: Path, output: Path) -> None:
    payload = json.loads(content_path.read_text(encoding='utf-8'))
    validate_content(payload)
    prs = Presentation()
    prs.slide_width = Inches(TOKENS['canvas']['width'])
    prs.slide_height = Inches(TOKENS['canvas']['height'])
    prs.core_properties.title = payload.get('meta', {}).get('title', 'Reference-style editable starter')
    prs.core_properties.subject = 'Reusable slide style derived from a user-provided reference'
    prs.core_properties.author = ''
    prs.core_properties.keywords = 'reference-slide-style; editable; starter; 16:9'
    set_theme(prs)
    for index, d in enumerate(payload['slides'], 1):
        BUILDERS[d['layout']](prs, payload.get('meta', {}), d, index)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    # Basic structural check; visual proofing remains necessary.
    with ZipFile(output) as z:
        corrupt = z.testzip()
        if corrupt:
            raise ValueError(f'ZIP validation failed: {corrupt}')
    check = Presentation(str(output))
    if len(check.slides) != len(payload['slides']):
        raise ValueError('Slide-count verification failed.')
    print(f'Created {output} ({len(check.slides)} editable slides).')


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--content', type=Path, default=ROOT / 'assets' / 'starter-content.json')
    parser.add_argument('--output', type=Path, default=ROOT / 'assets' / 'starter-layouts.pptx')
    args = parser.parse_args()
    try:
        make_deck(args.content, args.output)
    except (ValueError, KeyError, OSError, json.JSONDecodeError) as exc:
        parser.exit(2, f'Build failed: {exc}\n')


if __name__ == '__main__':
    main()
