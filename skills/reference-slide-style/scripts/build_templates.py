#!/usr/bin/env python3
"""Build bilingual blank PowerPoint templates with genuine named slide layouts.

Uses the bundled v1 starter only as a geometry blueprint. No research content is
copied. The four new layouts are Single content, Two columns, Decision, and
Thank you (Decision reuses the source's comparison geometry).

Run: python scripts/build_templates.py [--output-dir PATH]
The script does not render files; visual validation is a separate required step.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlideLayoutPart
from pptx.util import Inches, Pt
import build_deck as b

ROOT = Path(__file__).resolve().parents[1]
NS = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
LAYOUTS = [
 ('cover',0,'Cover','封面','1'),
 ('agenda',1,'Agenda','目录','2'),
 ('single',None,'Single content','单栏内容','v2 extension / v2 扩展'),
 ('two',None,'Two columns','双栏内容','v2 extension / v2 扩展'),
 ('overview',2,'Six-card overview','六卡总览','3, 15'),
 ('comparison',3,'Three-column comparison','三栏比较','4, 6, 8, 10'),
 ('architecture',4,'Workflow + architecture','流程与架构','5, 7, 9, 11, 12, 14'),
 ('table',5,'Evidence table','证据表格','13'),
 ('platform',6,'Shared platform','共享平台','16'),
 ('process',7,'Three-stage process','三阶段流程','17'),
 ('feedback',8,'Feedback loop','反馈闭环','18'),
 ('roadmap',9,'Roadmap + gates','路线图与门槛','19'),
 ('decision',3,'Closing decision','总结与决策','20'),
 ('thanks',None,'Thank you','致谢','v2 extension / v2 扩展'),
]

# Short field prompts, not sample business content. The prompts live in layouts;
# blank slides contain empty editable placeholders, except the final salutation.
PROMPTS = {
 'HEADLINE':('Headline','页面标题'), 'SUBTITLE':('Subtitle / scope','副标题／范围'),
 'FOOTER_LABEL':('Project / section','项目／章节'), 'BRAND_PLACEHOLDER':('Brand','品牌'),
 'PAGE_NUMBER':('Page','页码'), 'SOURCE_NOTE':('Source / date / caveat','来源／日期／说明'),
 'HERO_PLACEHOLDER_LABEL':('Image','图片'), 'COVER_KICKER':('Subject','主题'),
 'COVER_TITLE':('Presentation title','演示标题'), 'COVER_SUBTITLE':('Audience / subtitle','对象／副标题'),
 'COVER_META':('Presenter / team / date','汇报人／团队／日期'),
 'TAKEAWAY':('Key takeaway','核心结论'), 'SYSTEM_TITLE':('System name','系统名称'),
 'PLATFORM_BAND_TEXT':('Platform / contract','平台／协同约定'),
 'DEPENDENCY_TEXT':('Dependencies','依赖条件'),
 'FOUNDATION_TITLE':('Shared platform','共享平台'),
 'FOUNDATION_NOTE':('Responsibility / boundary','职责／边界'),
 'INPUT_OUTPUT_TEXT':('Inputs\nOutputs','输入\n输出'),
 'PROCESS_SEQUENCE':('Sequence / checkpoints','流程顺序／检查点'),
 'FEEDBACK_INPUT_TITLE':('Feedback inputs','反馈入口'),
 'FEEDBACK_LEDGER_TEXT':('Issue log / traceability','问题记录／追踪标识'),
 'FEEDBACK_RETURN_TEXT':('Return / validate / release','回访／验收／发布'),
 'INTERNAL_TITLE':('Internal validation','内部验证'), 'INTERNAL_BODY':('Measures / evidence','指标／证据'),
 'EXTERNAL_TITLE':('External validation','外部验证'), 'EXTERNAL_BODY':('Measures / evidence','指标／证据'),
 'IMPROVEMENT_TEXT':('Feedback / improvement','反馈／改进'),
 'CONTENT_TITLE':('Section heading','内容小标题'), 'CONTENT_BODY':('Content','正文内容'),
 'LEFT_TITLE':('Left heading','左栏标题'), 'LEFT_BODY':('Left content','左栏内容'),
 'RIGHT_TITLE':('Right heading','右栏标题'), 'RIGHT_BODY':('Right content','右栏内容'),
 'THANKS_TITLE':('Thank you','谢谢聆听'), 'THANKS_SUBTITLE':('Questions & discussion','欢迎交流与提问'),
 'THANKS_META':('Presenter / team','汇报人／团队'), 'THANKS_CONTACT':('Contact / date','联系方式／日期'),
}
PREFIXES = {
 'AGENDA_NUMBER':('Number','序号'), 'AGENDA_TITLE':('Section','章节'), 'AGENDA_DETAIL':('Scope / pages','范围／页码'),
 'CARD_NUMBER':('Number','序号'), 'CARD_LABEL':('Category','类别'), 'CARD_TITLE':('Card title','卡片标题'),
 'CARD_BODY':('Content','正文'), 'CARD_NOTE':('Note / boundary','说明／边界'),
 'WORKFLOW_TEXT':('Step','步骤'), 'LAYER_TITLE':('Layer title','层级标题'), 'LAYER_BODY':('Layer detail','层级说明'),
 'CALLOUT_TITLE':('Callout heading','说明标题'), 'CALLOUT_BODY':('Callout detail','详细说明'),
 'APPLICATION_TITLE':('Application','应用'), 'APPLICATION_BODY':('Task / output','任务／输出'),
 'CAPABILITY_TEXT':('Capability','能力'), 'PROCESS_LABEL':('Stage','阶段'),
 'PROCESS_TITLE':('Stage title','阶段标题'), 'PROCESS_BODY':('Tasks / outputs','任务／产出'),
 'INPUT_LABEL':('Input','入口'), 'INPUT_DETAIL':('Detail','说明'),
 'FEEDBACK_STAGE_TITLE':('Stage','环节'), 'FEEDBACK_STAGE_BODY':('Details','说明'),
}
PURPOSES = {
 'cover':('Presentation opening; keep the optional image in the right-hand picture placeholder.','演示开场；可选图片放入右侧图片占位符。'),
 'agenda':('Three sections; each band holds a number, heading and scope.','三个章节；每条横带包含序号、章节标题和范围。'),
 'single':('One open content area for a key idea, image or chart.','一个主要内容区，适合核心观点、图片或图表。'),
 'two':('Two parallel areas with a shared takeaway.','两个并列内容区，底部保留共同结论。'),
 'overview':('Six parallel cards; do not invent content to fill unused positions.','六张并列卡片；不得为了填满位置而编造内容。'),
 'comparison':('Three comparable subjects: category, title, evidence and caveat.','三个可比对象：类别、标题、证据和限制。'),
 'architecture':('Seven workflow chips, four system layers and four callouts.','七个流程步骤、四层系统架构和四组说明。'),
 'table':('A five-column native table with ten body rows. Duplicate this prepared slide to retain its table styling. A newly inserted layout provides a table placeholder, not a prefilled styled grid.','五列原生表格，预设十行正文。复制此预制页面可保留表格样式；从版式中新建页面只提供表格占位符，不自动生成带样式的网格。'),
 'platform':('Four application cards above five shared capabilities.','四个应用位于五项共享能力上方。'),
 'process':('Three sequential stages plus a checkpoint line.','三个连续阶段与一条检查点说明。'),
 'feedback':('A left-hand return loop and two right-hand validation areas.','左侧反馈回路，右侧两组验证说明。'),
 'roadmap':('Three periods: phase, deliverables and exit gate.','三个时期：阶段、交付物和进入下一阶段的门槛。'),
 'decision':('Three closing areas: priorities, support and next decision.','三个收束区域：重点、支持和下一步决策。'),
 'thanks':('Final acknowledgement slide; only the salutation and invitation are prefilled. Optional presenter and contact placeholders remain empty.','最终致谢页；仅预置致谢标题与交流邀请，汇报人及联系方式占位符留空。'),
}


def prompt_for(name: str, lang: str, layout_key: str) -> str:
    j = 1 if lang == 'zh-CN' else 0
    if name in PROMPTS:
        return PROMPTS[name][j]
    for pref, vals in PREFIXES.items():
        if name.startswith(pref+'_'):
            idx=name.rsplit('_',1)[-1]
            if layout_key == 'roadmap' and pref in ('CARD_LABEL','CARD_TITLE','CARD_BODY','CARD_NOTE'):
                vals={'CARD_LABEL':('Period','时期'),'CARD_TITLE':('Phase','阶段'),
                      'CARD_BODY':('Deliverables','交付物'),'CARD_NOTE':('Exit gate','进入下一阶段的门槛')}[pref]
            if layout_key == 'decision' and pref == 'CARD_TITLE':
                return [('Priorities','重点'),('Support','支持'),('Next decision','下一步决策')][int(idx)-1][j]
            # Keep body prompts short; numbering is only useful on repeated headings.
            return vals[j] + (' '+idx if pref not in ('CARD_BODY','CARD_NOTE','LAYER_BODY','CALLOUT_BODY','INPUT_DETAIL','FEEDBACK_STAGE_BODY','APPLICATION_BODY','PROCESS_BODY') else '')
    return '内容' if j else 'Content'


def first_run_style(shape):
    found=shape._element.xpath('.//a:rPr')
    if found:
        return deepcopy(found[0])
    el=OxmlElement('a:rPr');el.set('sz','1700');return el


def replace_text(shape, text_value, lang='en', field=True):
    """Preserve typography in paragraph defaults as well as typed runs."""
    tf=shape.text_frame
    proto=first_run_style(shape)
    code='zh-CN' if lang=='zh-CN' else 'en-US'
    proto.set('lang',code)
    for tag, face in [('a:latin','Arial'),('a:ea','Noto Sans CJK SC')]:
        for el in proto.findall(qn(tag)): proto.remove(el)
        el=OxmlElement(tag);el.set('typeface',face);proto.append(el)
    pstyle=deepcopy(tf.paragraphs[0]._p.pPr) if tf.paragraphs[0]._p.pPr is not None else OxmlElement('a:pPr')
    for old in list(pstyle):
        if old.tag in [qn('a:defRPr'),qn('a:buChar'),qn('a:buAutoNum'),qn('a:buNone')]:pstyle.remove(old)
    no=OxmlElement('a:buNone');pstyle.append(no)
    default=deepcopy(proto);default.tag=qn('a:defRPr');pstyle.append(default)
    # Ensure all placeholder levels preserve the intended size and no bullets.
    lst=tf._txBody.find(qn('a:lstStyle'))
    if lst is None:
        lst=OxmlElement('a:lstStyle');tf._txBody.insert(1,lst)
    for el in list(lst):lst.remove(el)
    for level in range(1,10):
        lv=deepcopy(pstyle);lv.tag=qn(f'a:lvl{level}pPr');lst.append(lv)
    for p in list(tf._txBody.findall(qn('a:p'))): tf._txBody.remove(p)
    for line in text_value.split('\n'):
        p=OxmlElement('a:p');p.append(deepcopy(pstyle))
        if line:
            r=OxmlElement('a:r');r.append(deepcopy(proto));t=OxmlElement('a:t');t.text=line;r.append(t);p.append(r)
        end=deepcopy(proto);end.tag=qn('a:endParaRPr');p.append(end);tf._txBody.append(p)
    # Local invisible appearance prevents body-placeholder default outlines/fills.
    sp=shape._element.spPr
    if not any(e.tag in [qn('a:solidFill'),qn('a:noFill'),qn('a:gradFill')] for e in sp):sp.append(OxmlElement('a:noFill'))
    return proto


def new_template_slide(prs, key):
    s=b.chrome(prs, {}, {'layout':key,'title':'Title','subtitle':'Subtitle'},1,cover=(key=='thanks'))
    if key=='single':
        b.box(s,'CONTENT_PANEL',[.58,2.12,12.17,3.93],'pale_blue',True)
        b.text(s,'CONTENT_TITLE',[.85,2.40,11.63,.45],'Heading',24,'navy',True)
        b.text(s,'CONTENT_BODY',[.85,3.05,11.63,2.65],'Content',19,'slate')
        b.takeaway(s,'Key takeaway')
    elif key=='two':
        for x,pref,fill in [(.58,'LEFT','pale_blue'),(6.82,'RIGHT','pale_purple')]:
            b.box(s,pref+'_PANEL',[x,2.12,5.93,3.93],fill,True)
            b.box(s,pref+'_ACCENT',[x,2.12,5.93,.045],'blue' if pref=='LEFT' else 'purple')
            b.text(s,pref+'_TITLE',[x+.25,2.45,5.43,.55],'Heading',24,'navy',True)
            b.text(s,pref+'_BODY',[x+.25,3.23,5.43,2.43],'Content',19,'slate')
        b.takeaway(s,'Key takeaway')
    elif key=='thanks':
        b.box(s,'THANKS_PANEL',[.58,1.80,12.17,3.55],'pale_blue',True)
        for i,c in enumerate(['blue','purple','teal']):
            b.box(s,'THANKS_ACCENT_'+str(i+1),[5.53+i*.79,2.24,.70,.045],c)
        b.text(s,'THANKS_TITLE',[1.20,2.68,10.93,.92],'Thank you',52,'blue',True,'center','middle',1)
        b.text(s,'THANKS_SUBTITLE',[1.20,3.78,10.93,.48],'Questions & discussion',22,'slate',False,'center','middle',1)
        b.text(s,'THANKS_META',[1.20,5.78,10.93,.4],'Presenter / team',18,'navy',False,'center','middle',1)
        b.text(s,'THANKS_CONTACT',[1.20,6.27,10.93,.30],'Contact / date',14,'slate',False,'center','middle',1)
    return s


def set_theme_language(prs,lang):
    b.set_theme(prs)
    for part in prs.part.package.iter_parts():
        if part.partname.startswith('/ppt/theme/'):
            root=etree.fromstring(part.blob)
            root.set('name','Reference Slide Style v2 - '+lang)
            for f in root.findall('.//a:font[@script="Hans"]',NS): f.set('typeface','Noto Sans CJK SC')
            for f in root.findall('.//a:ea',NS):f.set('typeface','Noto Sans CJK SC')
            part._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)


def add_placeholder(element,idx,typ='body'):
    nv=element.find('p:nvSpPr/p:nvPr',NS)
    if nv is None:nv=element.find('p:nvGraphicFramePr/p:nvPr',NS)
    for old in list(nv):
        if old.tag==qn('p:ph'):nv.remove(old)
    ph=OxmlElement('p:ph');ph.set('type',typ);ph.set('idx',str(idx));ph.set('hasCustomPrompt','1');nv.append(ph)


def create_layout(prs, blueprint, key, number, name,lang):
    master=prs.slide_master
    xml=f'''<p:sldLayout {nsdecls('a','p','r')} type="cust" preserve="1"><p:cSld name="{number:02d} | {name}"><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>'''
    root=parse_xml(xml)
    part=SlideLayoutPart(PackURI(f'/ppt/slideLayouts/slideLayout{number+20}.xml'),CT.PML_SLIDE_LAYOUT,prs.part.package,root)
    part.relate_to(master.part,RT.SLIDE_MASTER)
    rid=master.part.relate_to(part,RT.SLIDE_LAYOUT)
    lid=OxmlElement('p:sldLayoutId');lid.set('id',str(2147484000+number));lid.set(qn('r:id'),rid);master._element.sldLayoutIdLst.append(lid)
    layout=part.slide_layout
    fields=[];idx=0;native_table=None
    for original in blueprint.shapes:
        if original.name=='HERO_PLACEHOLDER_LABEL':continue
        el=deepcopy(original._element)
        if original.has_table:
            native_table=deepcopy(el)
            # New slides have an actual table placeholder; the prepared exemplar
            # below contains a native styled table using the same placeholder id.
            dummy=b.text(blueprint,'TMP_TABLE_PH',[original.left/914400,original.top/914400,original.width/914400,original.height/914400],'Table',13.5)
            el=deepcopy(dummy._element);blueprint.shapes._spTree.remove(dummy._element)
            el.find('p:nvSpPr/p:cNvPr',NS).set('name','EVIDENCE_TABLE')
            add_placeholder(el,idx,'tbl');add_placeholder(native_table,idx,'tbl')
            fields.append({'name':'EVIDENCE_TABLE','idx':idx,'type':'tbl','prompt':'表格' if lang=='zh-CN' else 'Table'})
            idx+=1
        elif original.has_text_frame and original.text:
            add_placeholder(el,idx,'title' if original.name in ('HEADLINE','COVER_TITLE','THANKS_TITLE') else 'body')
            fields.append({'name':original.name,'idx':idx,'type':'text','prompt':prompt_for(original.name,lang,key)})
            idx+=1
        layout.shapes._spTree.append(el)
    if key=='cover':
        # The original full-band image area stays a pale, blank background. Restrict
        # insertion to the exposed right side to avoid covering the white title panel.
        dum=b.text(blueprint,'COVER_PICTURE',[8.74,.70,4.5933333333,4.60],'Image',17,'purple',align='center',anchor='middle')
        el=deepcopy(dum._element);blueprint.shapes._spTree.remove(dum._element)
        add_placeholder(el,idx,'pic');layout.shapes._spTree.append(el)
        fields.append({'name':'COVER_PICTURE','idx':idx,'type':'pic','prompt':'图片' if lang=='zh-CN' else 'Image'})
    # Enforce unique shape IDs in the layout; no relationships needed by its geometry.
    for sid,el in enumerate(layout._element.xpath('.//p:spTree/*/p:nvSpPr/p:cNvPr | .//p:spTree/*/p:nvCxnSpPr/p:cNvPr | .//p:spTree/*/p:nvGraphicFramePr/p:cNvPr'),2):el.set('id',str(sid))
    for shape in layout.shapes:
        if shape.has_text_frame and shape.is_placeholder:
            prompt=next(f['prompt'] for f in fields if f['idx']==shape.placeholder_format.idx)
            replace_text(shape,prompt,lang)
            if key=='thanks' and shape.name=='THANKS_TITLE' and lang=='zh-CN':
                for rp in shape._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):rp.set('sz','4600')
    return layout,fields,native_table


def add_prepared_slide(prs,layout,fields,table,lang,number,key,guide,name,origin):
    s=prs.slides.add_slide(layout)
    fieldmap={f['idx']:f for f in fields}
    # Copy the full placeholder geometry/styles, not only idx. This remains editable
    # and gives predictable behavior in renderers with limited inheritance support.
    for sh in list(s.shapes):s.shapes._spTree.remove(sh._element)
    for ph in layout.placeholders:
        f=fieldmap[ph.placeholder_format.idx]
        if f['type']=='tbl':
            el=deepcopy(table);s.shapes._spTree.append(el)
            sh=s.shapes[-1]
            for i,row in enumerate(sh.table.rows):
                for j,cell in enumerate(row.cells):
                    # No selective evidence/status highlighting in an empty table.
                    cell.fill.solid();cell.fill.fore_color.rgb=b.rgb('blue' if i==0 else ('pale_blue' if i%2 else 'white'))
                    val=(f'C{j+1}' if lang=='en' else f'列{j+1}') if (guide and i==0) else ('内容' if lang=='zh-CN' else 'Entry') if guide and i==1 else ''
                    # Native table cells: keep run and paragraph defaults when empty.
                    tf=cell.text_frame
                    old=next((deepcopy(r._r.rPr) for p in tf.paragraphs for r in p.runs if r._r.rPr is not None),OxmlElement('a:rPr'))
                    old.set('lang','zh-CN' if lang=='zh-CN' else 'en-US')
                    ea=OxmlElement('a:ea');ea.set('typeface','Noto Sans CJK SC');old.append(ea)
                    p=tf.paragraphs[0]._p
                    for pp in list(tf._txBody.findall(qn('a:p')))[1:]:tf._txBody.remove(pp)
                    for cc in list(p):p.remove(cc)
                    ppr=OxmlElement('a:pPr');ppr.set('algn','l');no=OxmlElement('a:buNone');ppr.append(no)
                    de=deepcopy(old);de.tag=qn('a:defRPr');ppr.append(de);p.append(ppr)
                    if val:
                        r=OxmlElement('a:r');r.append(deepcopy(old));t=OxmlElement('a:t');t.text=val;r.append(t);p.append(r)
                    end=deepcopy(old);end.tag=qn('a:endParaRPr');p.append(end)
            continue
        el=deepcopy(ph._element);s.shapes._spTree.append(el);sh=s.shapes[-1]
        if f['type']=='pic':
            # Empty picture placeholder is not filled with a raster sample.
            if sh.has_text_frame:replace_text(sh,'',lang)
            if guide:
                b.text(s,'GUIDE_IMAGE_LABEL',[8.95,2.45,4.0,.65],'图片占位' if lang=='zh-CN' else 'Image placeholder',17,'purple',align='center',anchor='middle',spacing=1.0)
            continue
        value=f['prompt'] if guide else ''
        if f['name'] in ('THANKS_TITLE','THANKS_SUBTITLE'):value=f['prompt']
        elif guide and f['name']=='HEADLINE':value=f'{number:02d}  |  {name}'
        elif guide and f['name']=='COVER_TITLE':value='封面版式' if lang=='zh-CN' else 'Cover layout'
        elif guide and f['name']=='PAGE_NUMBER':value=f'{number:02d}'
        elif guide and (f['name'].startswith('AGENDA_NUMBER_') or f['name'].startswith('CARD_NUMBER_')):value=f"{int(f['name'].rsplit('_',1)[-1]):02d}"
        replace_text(sh,value,lang)
    for sid,el in enumerate(s._element.xpath('.//p:spTree/*/p:nvSpPr/p:cNvPr | .//p:spTree/*/p:nvGraphicFramePr/p:cNvPr'),2):el.set('id',str(sid))
    # Empty body placeholders do not export prompt text. The final two strings are
    # actual editable content, intentionally retained as the thank-you template.
    j=1 if lang=='zh-CN' else 0
    if lang=='zh-CN':
        provenance='v2 风格扩展，原参考文件没有对应独立页面。' if 'v2 extension' in origin else f'参考演示第 {origin} 页。'
        note=f'版式 {number:02d}：{name}\n{PURPOSES[key][j]}\n来源：{provenance}\n'
        note+='底图和卡片位于自定义版式中，正常编辑时不会误移动。文字占位符的位置、字号和颜色已确定。\n空白模板中前十三页不含正文；占位提示只用于编辑，不应作为最终内容导出。\n复制对应页面并填入真实材料；不要继承原参考文件的品牌、日期和研究结论。\n封面图片区限制在右侧，以免覆盖左侧白色标题面板。\n中文字体为 Noto Sans CJK SC；未安装时请统一替换为微软雅黑或苹方简体，随后重新检查换行。\n生成成稿后逐页检查；致谢页通常放在正文末尾，附录之前。\n'
    else:
        provenance='New v2 style-aligned extension; not a standalone layout in the source.' if 'v2 extension' in origin else f'Source slide(s) {origin}.'
        note=f'LAYOUT {number:02d}: {name}\n{PURPOSES[key][j]}\nOrigin: {provenance}\n'
        note+='Background geometry lives in a named custom layout; normal editing does not move it. Text placeholders have predetermined geometry and type.\nThe first thirteen blank slides contain no body text. Layout prompts are editing aids, not presentation content.\nDuplicate the appropriate slide and fill it with supplied facts. Do not inherit source research, branding or dates.\nThe cover picture placeholder is restricted to the right side to keep the white title panel visible.\nChinese font: Noto Sans CJK SC; replace consistently with Microsoft YaHei or PingFang SC when unavailable, then inspect wrapping.\nReview every completed slide. Place the thank-you slide at the end of the main presentation, before an optional appendix.\n'
    note+=('\n【版式标注版】本文件显示字段名称，供选择版式；正式制作请使用空白模板。' if guide and j else '\nLABELLED GUIDE: field names are visible for layout selection; use the blank template for production.' if guide else '')
    s.notes_slide.notes_text_frame.text=note
    return s


def save_potx(pptx_path,potx_path):
    with ZipFile(pptx_path) as src,ZipFile(potx_path,'w',ZIP_DEFLATED) as out:
        for inf in src.infolist():
            data=src.read(inf.filename)
            if inf.filename=='[Content_Types].xml':
                root=etree.fromstring(data)
                for e in root:
                    if e.get('PartName')=='/ppt/presentation.xml':e.set('ContentType','application/vnd.openxmlformats-officedocument.presentationml.template.main+xml')
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            out.writestr(inf,data)


def load_blueprint():
    """Use the optional bundled deck, or rebuild it from tracked JSON sources."""
    bundled = ROOT / 'assets' / 'starter-layouts.pptx'
    if bundled.is_file():
        return Presentation(bundled)
    with TemporaryDirectory(prefix='reference-slide-style-') as tmp:
        generated = Path(tmp) / 'starter-layouts.pptx'
        b.make_deck(ROOT / 'assets' / 'starter-content.json', generated)
        return Presentation(generated)


def make_all(outdir):
    source=load_blueprint()
    # New general-purpose and thank-you layouts use the same measurement helpers.
    extra=Presentation();extra.slide_width=Inches(13.3333333333);extra.slide_height=Inches(7.5)
    newslides={key:new_template_slide(extra,key) for key in ('single','two','thanks')}
    for lang in ('en','zh-CN'):
        for guide in (False,True):
            prs=Presentation();prs.slide_width=Inches(13.3333333333);prs.slide_height=Inches(7.5)
            for sh in list(prs.slide_master.shapes):prs.slide_master.shapes._spTree.remove(sh._element)
            oldlayouts=list(prs.slide_layouts)
            set_theme_language(prs,lang)
            prs.core_properties.title=('参考风格空白模板' if lang=='zh-CN' else 'Reference style blank templates')+(' / Layout guide' if guide else '')
            prs.core_properties.subject='14 reusable fixed layouts; blank content; bilingual style kit v2'
            prs.core_properties.author='';prs.core_properties.keywords='reference-slide-style; v2; blank; custom layouts; bilingual'
            for number,(key,srcidx,en,zh,origin) in enumerate(LAYOUTS,1):
                bp=source.slides[srcidx] if srcidx is not None else newslides[key]
                name=zh if lang=='zh-CN' else en
                layout,fields,table=create_layout(prs,bp,key,number,name,lang)
                add_prepared_slide(prs,layout,fields,table,lang,number,key,guide,name,origin)
            for lo in oldlayouts:prs.slide_layouts.remove(lo)
            filename=('layout-guide-' if guide else 'blank-layouts-')+lang+'.pptx'
            path=outdir/filename;prs.save(path)
            with ZipFile(path) as z:
                if z.testzip():raise RuntimeError('Corrupt archive: '+str(path))
            check=Presentation(path)
            assert len(check.slides)==14 and len(check.slide_layouts)==14
            if not guide:save_potx(path,outdir/f'reference-style-{lang}.potx')
            print(path)
    (outdir/'layout-index.json').write_text(json.dumps([{'number':i,'key':v[0],'name_en':v[2],'name_zh_CN':v[3],'source_slides':v[4]} for i,v in enumerate(LAYOUTS,1)],ensure_ascii=False,indent=2),encoding='utf-8')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,default=ROOT/'templates')
    a=parser.parse_args();a.output_dir.mkdir(parents=True,exist_ok=True)
    try:make_all(a.output_dir)
    except (ValueError,OSError,KeyError,AssertionError) as exc:parser.exit(2,f'Template build failed: {exc}\n')

if __name__=='__main__':main()
