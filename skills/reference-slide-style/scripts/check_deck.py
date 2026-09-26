#!/usr/bin/env python3
"""Structural QA for presentations built with reference-slide-style.

Does not measure actual rendered text, verify evidence, or prove accessibility.
Render and inspect every slide before delivering a final presentation.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
TOKENS = json.loads((ROOT / 'assets' / 'design-tokens.json').read_text(encoding='utf-8'))
INCH = 914400


def check(path: Path, allow_placeholders: bool = False) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    with ZipFile(path) as z:
        corrupt = z.testzip()
        if corrupt:
            errors.append(f'Corrupt ZIP member: {corrupt}')
        font_files = [name for name in z.namelist() if name.startswith('ppt/fonts/')]
        if font_files:
            warnings.append('Embedded font files are present; check distribution permissions.')
    prs = Presentation(str(path))
    width, height = prs.slide_width / INCH, prs.slide_height / INCH
    if abs(width - 13.3333333333) > .002 or abs(height - 7.5) > .002:
        errors.append(f'Expected 13.333333 x 7.5 inches; got {width:.6f} x {height:.6f}.')
    shape_count = 0
    table_count = 0
    text_count = 0
    placeholders = []
    for n, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        names = Counter(sh.name for sh in shapes)
        duplicates = [key for key, count in names.items() if count > 1]
        if duplicates:
            warnings.append(f'Slide {n}: duplicate shape names: {duplicates}')
        by_name = {sh.name: sh for sh in shapes}
        for role, token in [('FOOTER_RULE', 'footer_rule'), ('FOOTER_LABEL', 'footer_label'),
                            ('PAGE_NUMBER', 'page')]:
            if role not in by_name:
                warnings.append(f'Slide {n}: no named {role}; check its placement manually.')
                continue
            sh = by_name[role]
            observed = [v / INCH for v in (sh.left, sh.top, sh.width, sh.height)]
            if any(abs(a-b) > .025 for a,b in zip(observed, TOKENS['chrome'][token])):
                errors.append(f'Slide {n}: {role} differs from the expected frame.')
        if 'PAGE_NUMBER' in by_name and by_name['PAGE_NUMBER'].text.strip() != f'{n:02d}':
            errors.append(f'Slide {n}: page number is not {n:02d}.')
        if 'HEADLINE' in by_name:
            sh = by_name['HEADLINE']
            if any(abs(a/INCH-b) > .025 for a,b in zip((sh.left,sh.top,sh.width,sh.height), TOKENS['chrome']['headline'])):
                errors.append(f'Slide {n}: headline frame differs from the style.')
        for sh in shapes:
            shape_count += 1
            x,y,w,h = [v / INCH for v in (sh.left,sh.top,sh.width,sh.height)]
            if x < -.01 or y < -.01 or x+w > width+.01 or y+h > height+.01:
                errors.append(f'Slide {n}: {sh.name} extends beyond the slide canvas.')
            if sh._element.xpath('.//a:outerShdw | .//a:innerShdw'):
                errors.append(f'Slide {n}: {sh.name} contains an explicit shadow.')
            texts = []
            if sh.has_text_frame:
                texts.append(sh.text)
                if sh.text.strip():
                    text_count += 1
            if sh.has_table:
                table_count += 1
                tbl = sh.table
                tw = sum(col.width for col in tbl.columns) / INCH
                th = sum(row.height for row in tbl.rows) / INCH
                if abs(tw-w) > .025 or abs(th-h) > .025:
                    errors.append(f'Slide {n}: table frame does not match its grid dimensions.')
                texts.extend(cell.text for row in tbl.rows for cell in row.cells)
            for value in texts:
                if '\ufffd' in value:
                    warnings.append(f'Slide {n}: {sh.name} has a Unicode replacement character.')
                if re.search(r'\[[^\[\]\n]{1,120}\]', value):
                    placeholders.append(f'Slide {n}: {sh.name}')
    if placeholders and not allow_placeholders:
        errors.append(f'{len(placeholders)} objects contain bracketed placeholder-like text; replace or review it. '
                      'Use --allow-placeholders only for a reusable starter or reviewed intentional brackets.')
    if not len(prs.slides):
        errors.append('The presentation contains no slides.')
    return {
        'file': str(path), 'slides': len(prs.slides),
        'width_inches': width, 'height_inches': height,
        'shapes': shape_count, 'text_objects': text_count, 'native_tables': table_count,
        'placeholder_objects': len(placeholders), 'allow_placeholders': allow_placeholders,
        'errors': errors, 'warnings': warnings,
        'result': 'PASS' if not errors else 'FAIL',
        'limitations': ['No rendered-text fit check.', 'No factual or citation verification.',
                        'No cross-application or target-device font guarantee.'],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--allow-placeholders', action='store_true')
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    try:
        report = check(args.input, args.allow_placeholders)
    except Exception as exc:
        parser.exit(2, f'Unable to check presentation: {exc}\n')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    parser.exit(0 if report['result'] == 'PASS' else 1)


if __name__ == '__main__':
    main()
