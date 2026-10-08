#!/usr/bin/env python3
"""Copy selected unique exemplars into a working PPTX; requires python-pptx."""

import argparse
from pathlib import Path

from pptx import Presentation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input", type=Path,
        default=Path(__file__).resolve().parents[1] / "assets" / "templates.pptx",
        help="Template library (defaults to the bundled deck)",
    )
    parser.add_argument("--slides", required=True, help="One-based positions, e.g. 9,6,7")
    parser.add_argument("--output", type=Path, required=True, help="New working PPTX")
    args = parser.parse_args()
    try:
        positions = [int(value.strip()) for value in args.slides.split(",")]
    except ValueError:
        parser.error("--slides must contain comma-separated integers")
    if len(set(positions)) != len(positions):
        parser.error("Select unique exemplars; duplicate pages with the authoring tool")
    if args.output.exists():
        parser.error("Output already exists; choose a new file to preserve existing work")

    deck = Presentation(args.input)
    if any(position < 1 or position > len(deck.slides) for position in positions):
        parser.error(f"Slide positions must be between 1 and {len(deck.slides)}")
    ids = list(deck.slides._sldIdLst)
    selected = set(positions)
    for position, slide_id in enumerate(ids, 1):
        if position not in selected:
            deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    for position in positions:
        deck.slides._sldIdLst.append(ids[position - 1])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(args.output)
    print(f"Selected {positions} into {args.output.resolve()}")


if __name__ == "__main__":
    main()
